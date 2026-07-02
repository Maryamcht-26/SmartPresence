from openpyxl import load_workbook
from sqlalchemy.orm import Session
from sqlalchemy import text
from datetime import time, datetime
import re


def parse_time(value):
    if isinstance(value, time):
        return value
    if isinstance(value, str):
        value = value.strip()
        match = re.match(r'(\d+)[Hh:](\d+)', value)
        if match:
            return time(int(match.group(1)), int(match.group(2)))
        return datetime.strptime(value, "%H:%M").time()
    return value


def parse_time_range(text_val):
    if not text_val:
        return None, None
    match = re.match(r'(\d+[Hh:]\d+)\s*[àa]\s*(\d+[Hh:]\d+)', str(text_val), re.IGNORECASE)
    if match:
        return parse_time(match.group(1)), parse_time(match.group(2))
    return None, None


NATURE_MAP = {
    "cours":        "Cours",
    "td":           "TD",
    "tp":           "TP",
    "cours/td":     "Cours/TD",
    "cours/tp":     "Cours/TP",
    "cours/td/tp":  "Cours/TD/TP",
}

def normalize_nature(value):
    if not value:
        return "Cours"
    clean = re.sub(r'\s*→.*$', '', str(value), flags=re.IGNORECASE)
    clean = re.sub(r'\s*/?\s*GR\w*', '', clean, flags=re.IGNORECASE)
    clean = re.sub(r'\s*/?\s*Groupe\s*\d+', '', clean, flags=re.IGNORECASE)
    clean = clean.strip().strip('/')
    return NATURE_MAP.get(clean.lower(), "Cours")


def is_merged_across(sheet, row, col_start, col_end):
    """
    Vérifie si la cellule à (row, col_start) fait partie d'une fusion
    qui couvre AU MOINS col_start jusqu'à col_end sur cette ligne.

    Si oui => la case "Matière" n'est pas divisée par groupe
    => il ne faut affecter AUCUN groupe (groupe_id = None = tout le niveau).
    """
    for merged_range in sheet.merged_cells.ranges:
        if (merged_range.min_row <= row <= merged_range.max_row
                and merged_range.min_col <= col_start
                and merged_range.max_col >= col_end):
            return True
    return False


def get_or_create_matiere(db: Session, nom: str):
    mat = db.execute(
        text("SELECT id FROM matiere WHERE nom = :nom"),
        {"nom": nom}
    ).fetchone()

    if mat:
        return mat[0]

    result = db.execute(
        text("INSERT INTO matiere (nom) VALUES (:nom) RETURNING id"),
        {"nom": nom}
    )
    return result.fetchone()[0]



def get_or_create_groupe(db: Session, nom: str, niveau_id: int):
    if not nom:
        return None

    grp = db.execute(
        text("""
            SELECT id FROM groupe
            WHERE nom=:nom AND niveau_id=:cid
        """),
        {"nom": nom, "cid": niveau_id}
    ).fetchone()

    if grp:
        return grp[0]

    result = db.execute(
        text("""
            INSERT INTO groupe (nom, niveau_id)
            VALUES (:nom, :cid)
            RETURNING id
        """),
        {"nom": nom, "cid": niveau_id}
    )
    return result.fetchone()[0]


from app.core.security import hash_password

def get_or_create_prof(db: Session, nom_complet: str):
    if not nom_complet:
        return None

    parts = nom_complet.strip().split()
    nom = parts[0]
    prenom = " ".join(parts[1:]) if len(parts) > 1 else ""

    prof = db.execute(
        text("SELECT id FROM prof WHERE nom=:nom AND prenom=:prenom"),
        {"nom": nom, "prenom": prenom}
    ).fetchone()

    if prof:
        return prof[0]

    year = datetime.now().year
    nom_clean = nom.lower().replace(" ", "")
    prenom_clean = prenom.lower().replace(" ", "")

    identity = ".".join(p for p in [prenom_clean, nom_clean] if p)
    email = f"{identity}@ump.ac.ma"
    temp_password = f"{nom_clean}{prenom_clean}{year}"

    # Use centralized security helper
    hashed = hash_password(temp_password)

    # Second check: check if the generated email already exists for another name
    existing_email = db.execute(
        text("SELECT id FROM prof WHERE email=:email"),
        {"email": email}
    ).fetchone()
    if existing_email:
        return existing_email[0]

    result = db.execute(
        text("""
            INSERT INTO prof (nom, prenom, email, password_hash, must_change_password)
            VALUES (:nom, :prenom, :email, :hash, true)
            RETURNING id
        """),
        {"nom": nom, "prenom": prenom, "email": email, "hash": hashed}
    )
    return result.fetchone()[0]


def get_current_semestre_annee():
    """
    Détermine le semestre et l'année universitaire actuels en fonction de la date du jour.
    - Septembre à Décembre, ou Janvier => S1
    - Février à Juillet => S2
    - Année universitaire : commence en septembre, se termine en juillet/août de l'année suivante.
    """
    now = datetime.now()
    mois = now.month
    annee_actuelle = now.year

    if mois >= 9 or mois == 1:
        semestre = "S1"
    else:
        semestre = "S2"

    if mois >= 9:
        # Septembre à décembre : l'année universitaire commence cette année
        annee = f"{annee_actuelle}-{annee_actuelle + 1}"
    else:
        # Janvier à juillet/août : l'année universitaire a commencé l'an dernier
        annee = f"{annee_actuelle - 1}-{annee_actuelle}"

    return semestre, annee


def get_or_create_creneau(db, jour, h_debut, h_fin, niveau_id, semestre, annee):
    # Fallback dynamique si les métadonnées sont absentes de l'Excel
    if not semestre or not annee:
        default_semestre, default_annee = get_current_semestre_annee()
        if not semestre:
            semestre = default_semestre
        if not annee:
            annee = default_annee

    res = db.execute(
        text("""
            SELECT id FROM creneau
            WHERE jour=:j AND heure_debut=:hd AND niveau_id=:cid
            AND semestre=:sem AND annee_universitaire=:annee
        """),
        {"j": jour, "hd": h_debut, "cid": niveau_id, "sem": semestre, "annee": annee}
    ).fetchone()

    if res:
        return res[0]

    result = db.execute(
        text("""
            INSERT INTO creneau
            (jour, heure_debut, heure_fin, semestre, annee_universitaire, niveau_id)
            VALUES (:j, :hd, :hf, :sem, :annee, :cid)
            RETURNING id
        """),
        {"j": jour, "hd": h_debut, "hf": h_fin, "cid": niveau_id, "sem": semestre, "annee": annee}
    )
    return result.fetchone()[0]

def insert_creneau_groupe(db, creneau_id, groupe_id, matiere_id, prof_id, salle, nature):
    db.execute(
        text("""
            INSERT INTO creneau_groupe
            (creneau_id, groupe_id, matiere_id, prof_id, salle, nature)
            VALUES (:cr, :grp, :mat, :prof, :salle, :nature)
            ON CONFLICT DO NOTHING
        """),
        {
            "cr": creneau_id,
            "grp": groupe_id,
            "mat": matiere_id,
            "prof": prof_id,
            "salle": salle,
            "nature": nature
        }
    )


def detect_format(sheet):
    row2 = [sheet.cell(2, c).value for c in range(1, 5)]
    if any(str(v).strip().lower() == 'jour' for v in row2 if v):
        return 'tabular'
    return 'matrix'



def parse_tabular(sheet, db: Session, niveau_id: int, semestre, annee):
    for row in sheet.iter_rows(min_row=3, values_only=True):
        jour, h_debut, h_fin, groupe, matiere, nature, prof, salle = row

        if not jour or not matiere or not prof:
            continue

        jour = str(jour).strip().upper()
        h_debut = parse_time(h_debut)
        h_fin = parse_time(h_fin)

        matiere_id = get_or_create_matiere(db, matiere.strip())
        groupe_id = get_or_create_groupe(db, groupe, niveau_id)
        prof_id = get_or_create_prof(db, prof)

        creneau_id = get_or_create_creneau(
            db, jour, h_debut, h_fin, niveau_id, semestre, annee
        )
        insert_creneau_groupe(db, creneau_id, groupe_id, matiere_id, prof_id, salle, nature)



JOURS = {"LUNDI", "MARDI", "MERCREDI", "JEUDI", "VENDREDI", "SAMEDI", "DIMANCHE"}
FIELDS = ["Matière", "Nature Ens", "Professeur", "Charge", "Semaines", "Nbre semaine", "Salle"]


def parse_matrix(sheet, db: Session, niveau_id: int, semestre, annee):
    max_row = sheet.max_row
    max_col = sheet.max_column

    TIME_ROW, GROUP_ROW = 5, 6
    time_slots = []
    col = 3
    while col <= max_col:
        cell_val = sheet.cell(TIME_ROW, col).value
        if cell_val:
            h_debut, h_fin = parse_time_range(str(cell_val))
            if h_debut:
                time_slots.append((h_debut, h_fin, col))
        col += 1

    group_map = {}
    for (h_debut, h_fin, col_start) in time_slots:
        for offset in range(2):
            c = col_start + offset
            if c > max_col:
                break
            grp_val = sheet.cell(GROUP_ROW, c).value
            group_map[c] = str(grp_val).strip() if grp_val else None

    current_jour = None
    day_block = {}
    day_block_rows = {}  #mémorise le numéro de ligne réel de chaque champ pour ce jour

    def flush_block(jour, block, block_rows):
        # Ligne de référence pour détecter la fusion : la ligne "Matière" du bloc de ce jour
        matiere_row = block_rows.get("Matière")

        for (h_debut, h_fin, col_start) in time_slots:
            col_end = col_start + 1
            if col_end > max_col:
                col_end = col_start

            # Détection de fusion : si la cellule "Matière" est fusionnée sur
            # les 2 colonnes (Groupe1 + Groupe2) du créneau => case non divisée
            # => on n'affecte aucun groupe.
            fused = False
            if matiere_row:
                fused = is_merged_across(sheet, matiere_row, col_start, col_end)

            already_inserted = False  # évite la double insertion quand fused=True

            for offset in range(2):
                c = col_start + offset
                if c > max_col:
                    break

                matiere = block.get("Matière", {}).get(c)
                if not matiere:
                    continue
                matiere = str(matiere).strip()

                nature = block.get("Nature Ens", {}).get(c)
                nature = normalize_nature(nature)

                prof_raw = block.get("Professeur", {}).get(c)
                prof_raw = str(prof_raw).strip() if prof_raw else None
                groupe_nom = group_map.get(c)

                if not prof_raw:
                    print(f"Skipped: {jour} {h_debut} | {matiere} | groupe={groupe_nom} - pas de prof")
                    continue

                salle = block.get("Salle", {}).get(c)
                if not salle:
                    salle = block.get("Charge", {}).get(c)
                salle = str(salle).strip() if salle else None

                # RÈGLE CLÉ : si la case n'est pas divisée (fusionnée),
                # on n'affecte aucun groupe (groupe_id = None = tout le niveau),
                # et on n'insère qu'UNE seule fois pour ce créneau.
                if fused:
                    if already_inserted:
                        continue
                    groupe_nom_final = None
                    already_inserted = True
                else:
                    groupe_nom_final = groupe_nom

                matiere_id = get_or_create_matiere(db, matiere)
                groupe_id = get_or_create_groupe(db, groupe_nom_final, niveau_id) if groupe_nom_final else None
                prof_id = get_or_create_prof(db, prof_raw) if prof_raw else None

                creneau_id = get_or_create_creneau(
                    db, jour.upper(), h_debut, h_fin, niveau_id, semestre, annee
                )
                insert_creneau_groupe(db, creneau_id, groupe_id, matiere_id, prof_id, salle, nature)

    for r in range(1, max_row + 1):
        col_a = sheet.cell(r, 1).value
        col_b = sheet.cell(r, 2).value
        col_a_str = str(col_a).strip().upper() if col_a else ""
        col_b_str = str(col_b).strip() if col_b else ""

        if col_a_str in JOURS:
            if current_jour and day_block:
                flush_block(current_jour, day_block, day_block_rows)
            current_jour = col_a_str
            day_block = {f: {} for f in FIELDS}
            day_block_rows = {}  #reset pour le nouveau jour

        if current_jour and col_b_str in FIELDS:
            day_block_rows[col_b_str] = r  # mémorise la ligne réelle de ce champ
            for c in range(3, max_col + 1):
                val = sheet.cell(r, c).value
                if val is not None:
                    day_block[col_b_str][c] = val

    if current_jour and day_block:
        flush_block(current_jour, day_block, day_block_rows)




def parse_and_save(path, db: Session, niveau_id: int):
    wb = load_workbook(path)
    sheet = wb.active

    semestre, annee = parse_metadata(sheet)

    fmt = detect_format(sheet)

    if fmt == 'tabular':
        parse_tabular(sheet, db, niveau_id, semestre, annee)
    else:
        parse_matrix(sheet, db, niveau_id, semestre, annee)

    db.commit()


def parse_metadata(sheet):
    semestre = None
    annee = None

    for r in range(1, 15):
        col_a = str(sheet.cell(r, 1).value or "").strip().lower()
        col_c = str(sheet.cell(r, 3).value or "").strip()

        if "semestre" in col_a and col_c:
            semestre = col_c.upper()

        if "universitaire" in col_a and col_c:
            annee = col_c

    return semestre, annee